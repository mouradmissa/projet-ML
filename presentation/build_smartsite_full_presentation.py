# -*- coding: utf-8 -*-
"""
SmartSite — présentation PowerPoint (python-pptx).
Contenu métier : textes Introduction / Business Understanding fournis, intacts.
Données ML : issues des pipelines et scripts d'évaluation du dépôt.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

PRESENTATION_DIR = Path(__file__).resolve().parent
ML_ROOT = PRESENTATION_DIR.parent
OUT_PATH = PRESENTATION_DIR / "SmartSite_IA_Gestion_Chantier.pptx"

# Thème
C_BG = RGBColor(15, 23, 42)  # #0F172A
C_ACCENT = RGBColor(59, 130, 246)  # #3B82F6
C_WHITE = RGBColor(255, 255, 255)
C_MUTED = RGBColor(148, 163, 184)  # slate-400 sur fond sombre si besoin


def _set_slide_dark_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = C_BG


def _style_title_shape(title_shape, size_pt: int = 28) -> None:
    tf = title_shape.text_frame
    tf.word_wrap = True
    for p in tf.paragraphs:
        p.font.bold = True
        p.font.size = Pt(size_pt)
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.LEFT


def _add_title_only_slide(prs: Presentation, title: str) -> None:
    layout = prs.slide_layouts[5]  # Title only
    slide = prs.slides.add_slide(layout)
    _set_slide_dark_bg(slide)
    _style_title_shape(slide.shapes.title, 26)
    slide.shapes.title.text = title


def _add_bullets(
    prs: Presentation,
    title: str,
    bullets: list[str],
    *,
    title_size: int = 26,
    body_size: int = 15,
    max_width: float = 12.2,
) -> None:
    layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(layout)
    _set_slide_dark_bg(slide)
    title_sh = slide.shapes.title
    _style_title_shape(title_sh, title_size)
    title_sh.text = title

    left, top, w, h = Inches(0.45), Inches(1.15), Inches(max_width), Inches(5.9)
    box = slide.shapes.add_textbox(left, top, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(body_size)
        p.font.color.rgb = C_WHITE
        p.space_after = Pt(6)


def _add_two_column(
    prs: Presentation,
    title: str,
    left_title: str,
    left_bullets: list[str],
    right_title: str,
    right_bullets: list[str],
) -> None:
    layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(layout)
    _set_slide_dark_bg(slide)
    _style_title_shape(slide.shapes.title, 24)
    slide.shapes.title.text = title

    col_w = Inches(5.9)
    y0 = Inches(1.1)
    h = Inches(5.95)

    for x_off, sub_t, blist in (
        (Inches(0.4), left_title, left_bullets),
        (Inches(6.85), right_title, right_bullets),
    ):
        box = slide.shapes.add_textbox(x_off, y0, col_w, h)
        tf = box.text_frame
        tf.word_wrap = True
        p0 = tf.paragraphs[0]
        p0.text = sub_t
        p0.font.bold = True
        p0.font.size = Pt(17)
        p0.font.color.rgb = C_ACCENT
        p0.space_after = Pt(10)
        for line in blist:
            p = tf.add_paragraph()
            p.text = line
            p.level = 0
            p.font.size = Pt(14)
            p.font.color.rgb = C_WHITE
            p.space_after = Pt(5)


def _add_table_slide(
    prs: Presentation,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    *,
    col_widths: list[float] | None = None,
    header_pt: int = 11,
    cell_pt: int = 10,
) -> None:
    layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(layout)
    _set_slide_dark_bg(slide)
    _style_title_shape(slide.shapes.title, 22)
    slide.shapes.title.text = title

    nrows = len(rows) + 1
    ncols = len(headers)
    left = Inches(0.35)
    top = Inches(1.05)
    width = Inches(12.55)
    row_h = Inches(0.38)
    height = row_h * nrows

    table = slide.shapes.add_table(nrows, ncols, left, top, width, height).table
    if col_widths and len(col_widths) == ncols:
        for j, cw in enumerate(col_widths):
            table.columns[j].width = Inches(cw)

    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(header_pt)
            p.font.color.rgb = C_WHITE
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = table.cell(i + 1, j)
            c.text = val
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(cell_pt)
                p.font.color.rgb = C_WHITE


def _add_focus_slide(prs: Presentation, headline: str, sub: str | None = None) -> None:
    layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(layout)
    _set_slide_dark_bg(slide)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.5), Inches(2.2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = headline
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    if sub:
        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(20)
        p2.font.color.rgb = C_ACCENT
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(16)


def _cover_slide(prs: Presentation) -> None:
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    _set_slide_dark_bg(slide)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(12.3), Inches(3.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = "SmartSite"
    p.font.bold = True
    p.font.size = Pt(54)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = "IA pour la gestion de chantier"
    p2.font.size = Pt(26)
    p2.font.color.rgb = C_ACCENT
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(14)
    p3 = tf.add_paragraph()
    p3.text = "[Nom]  ·  [Date]"
    p3.font.size = Pt(16)
    p3.font.color.rgb = C_MUTED
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(36)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --- PAGE DE GARDE ---
    _cover_slide(prs)

    # --- INTRODUCTION (texte fourni, inchangé) ---
    _add_bullets(
        prs,
        "Contexte",
        [
            "Le secteur de la construction fait face à des défis majeurs : retards fréquents, dépassements budgétaires et manque de visibilité sur l'avancement des projets.",
            "Plus de 60 % des projets dépassent les délais ou le budget initialement prévus.",
        ],
        body_size=16,
    )

    _add_bullets(
        prs,
        "Solution SmartSite",
        [
            "SmartSite est une plateforme complète de gestion de chantier qui centralise :",
            "Le planning",
            "Les ressources",
            "La documentation",
            "La sécurité",
            "Le suivi financier",
            "Objectif : Permettre un pilotage en temps réel pour améliorer l'efficacité et la transparence des chantiers.",
        ],
    )

    _add_bullets(
        prs,
        "Problématique",
        [
            "Malgré les outils existants, les responsables de projets et les ingénieurs de chantier manquent d'outils prédictifs et intelligents pour :",
            "Anticiper les risques de retard ou de dépassement budgétaire",
            "Estimer précisément l'avancement réel des tâches",
            "Prioriser efficacement les tâches critiques",
        ],
    )

    _add_bullets(
        prs,
        "Conséquences",
        [
            "Prise de décision tardive",
            "Répartition sous-optimale des ressources",
            "Pertes financières et baisse de productivité",
        ],
    )

    _add_bullets(
        prs,
        "Objectif du Projet Machine Learning",
        [
            "Intégrer l'IA dans SmartSite afin de transformer les données opérationnelles en insights actionnables.",
            "Trois modèles de Machine Learning seront développés et comparés :",
            "Segmentation du Risque (Classification)",
            "Prédire le niveau de risque (Low / Medium / High) d'une tâche ou d'un projet.",
            "Prédiction de l'Avancement (Régression)",
            "Estimer le pourcentage d'avancement réel d'une tâche.",
            "Recommandation de Priorité (Classification)",
            "Attribuer automatiquement une priorité (Low / Medium / High) aux tâches.",
        ],
        body_size=14,
    )

    _add_bullets(
        prs,
        "Bénéfices Attendus",
        [
            "Améliorer le pilotage des chantiers",
            "Réduire les retards",
            "Optimiser l'allocation des ressources",
        ],
    )

    # --- BUSINESS UNDERSTANDING (texte fourni) ---
    _add_bullets(
        prs,
        "Compréhension du Contexte Métier",
        [
            "Problèmes Métiers Principaux",
            "Retards fréquents dans l'exécution des tâches et des projets",
            "Manque de visibilité en temps réel sur l'avancement réel des chantiers",
            "Suivi budgétaire complexe avec des dépassements fréquents",
            "Difficulté à prioriser et coordonner les ressources humaines et matérielles",
            "Faible automatisation de l'analyse des données terrain et des indicateurs de performance",
        ],
        body_size=14,
    )

    _add_table_slide(
        prs,
        "Alignement BOS - DSO",
        [
            "BOS (Objectif Métier)",
            "DSO (Objectif Data Science)",
            "Datasets / Variables Principales",
            "Type de Modèle",
        ],
        [
            [
                "Réduire les retards et anticiper les problèmes",
                "Prédire le niveau de risque d'une tâche/projet (Low / Medium / High)",
                "Risk_Level (cible), Progress, Budget, Actual Cost, Budget_Utilization, Planned_Duration_Days",
                "Classification",
            ],
            [
                "Améliorer le suivi et la visibilité de l'avancement",
                "Estimer précisément le pourcentage d'avancement réel des tâches",
                "Progress (cible), Start Date, End Date, Planned_Duration_Days, Hours Spent, Project Status, Task Status, Budget, Actual Cost",
                "Régression",
            ],
            [
                "Optimiser la priorisation des tâches",
                "Recommander automatiquement la priorité des tâches (Low / Medium / High)",
                "Priority (cible), Progress, Budget, Planned_Duration_Days, Risk_Level",
                "Classification",
            ],
        ],
        col_widths=[2.5, 2.85, 4.5, 1.7],
        header_pt=10,
        cell_pt=9,
    )

    _add_bullets(
        prs,
        "Parties Prenantes Clés",
        [
            "Project Manager : Pilotage global, planification et prise de décision",
            "Site Engineer : Suivi terrain, mise à jour de l'avancement et des incidents",
            "Finance : Suivi budgétaire, facturation et contrôle des coûts",
            "Client : Visibilité sur l'avancement et transparence du projet",
        ],
    )

    _add_bullets(
        prs,
        "Valeur Attendue du Projet Data Science",
        [
            "Anticiper les retards et les dépassements budgétaires",
            "Automatiser la priorisation des tâches critiques",
            "Produire des KPIs fiables et en temps réel",
            "Améliorer significativement la productivité et la rentabilité des projets",
        ],
    )

    # --- DATA : Segmentation du risque ---
    _add_focus_slide(prs, "Data — Segmentation du risque", "Variables · Prétraitement")

    _add_bullets(
        prs,
        "Segmentation du risque — Variables",
        [
            "Cible : Risk_Level (High / Low / Medium), encodée (LabelEncoder)",
            "Features : colonnes du CSV Project-Management-2-enriched.csv hors Risk_Level, après nettoyage",
            "Inclut variables numériques restantes et catégorielles (≤10 modalités) en one-hot",
        ],
    )

    _add_bullets(
        prs,
        "Segmentation du risque — Prétraitement",
        [
            "Suppression doublons ; colonnes constantes ; catégorielles >10 modalités ; corrélation >0,95",
            "Outliers continus : clip P1–P99 si hors bornes IQR (knn_risk_pipeline.py, risk_knn_svm_dt_evaluation.py)",
            "StandardScaler sur toutes les colonnes ; PCA 95 % variance si >20 features",
            "Split train/test : 80/20, stratify, random_state=42",
        ],
    )

    _add_bullets(
        prs,
        "Segmentation du risque — Screenshot",
        ["Insérer screenshot du preprocessing ici"],
    )

    # --- DATA : Avancement ---
    _add_focus_slide(prs, "Data — Prédiction de l'avancement", "Variables · Prétraitement")

    _add_bullets(
        prs,
        "Avancement — Variables",
        [
            "Cible : Progress (réelle dans [0, 1] dans les pipelines)",
            "Features : dérivées des dates (planned_duration_days, days_since_start, remaining_days),",
            "statuts ordinaux (Project Status, Task Status), numériques (Hours Spent, Budget, Actual Cost),",
            "one-hot sur catégorielles restantes ≤10 modalités ; Risk_Level et identifiants exclus (fuites)",
        ],
        body_size=13,
    )

    _add_bullets(
        prs,
        "Avancement — Prétraitement",
        [
            "Doublons supprimés ; NaN : drop lignes si <5 % de lignes concernées, sinon imputation (rf_progress_pipeline.py)",
            "IQR cap sur Hours Spent, Budget, Actual Cost",
            "Dates parsées (dayfirst) ; date de référence fixe 2026-04-12 pour jours écoulés / restants",
            "RF : pas de scaling ; LR : StandardScaler dans un Pipeline (lr_progress_pipeline.py)",
            "Split 80/20, random_state=42",
        ],
        body_size=13,
    )

    _add_bullets(
        prs,
        "Avancement — Screenshot",
        ["Insérer screenshot du preprocessing ici"],
    )

    # --- DATA : Priorité ---
    _add_focus_slide(prs, "Data — Recommandation de priorité", "Variables · Prétraitement")

    _add_bullets(
        prs,
        "Priorité — Variables",
        [
            "Cible : Priority (Low / Medium / High), encodage ordinal 0/1/2 dans les pipelines",
            "Features utilisées dans rf_priority_pipeline.py et knn_priority_pipeline.py :",
            "Progress, Budget, Planned_Duration_Days",
            "Note : le tableau BOS cite Risk_Level ; les modèles RF/KNN du dépôt n'intègrent pas Risk_Level dans X.",
        ],
        body_size=14,
    )

    _add_bullets(
        prs,
        "Priorité — Prétraitement",
        [
            "Doublons supprimés ; imputation médiane / mode si NaN (pipelines)",
            "Outliers IQR : plafonnement sur Budget et Planned_Duration_Days",
            "Split 80/20 stratifié sur la cible, random_state=42",
            "KNN : StandardScaler (obligatoire) — évaluation_priorite_rf_knn.py",
        ],
    )

    _add_bullets(
        prs,
        "Priorité — Screenshot",
        ["Insérer screenshot du preprocessing ici"],
    )

    # --- MODÉLISATION : Risque ---
    _add_focus_slide(prs, "Modélisation — Segmentation du risque", None)

    _add_bullets(
        prs,
        "Risque — Modèles utilisés",
        [
            "KNeighborsClassifier (k choisi 1–10 sur accuracy test, risk_knn_svm_dt_evaluation.py)",
            "SVC noyau RBF, probability=True, GridSearchCV C / gamma, CV stratifiée 5-fold",
            "DecisionTreeClassifier, GridSearchCV puis élagage par coût-complexité (ccp_alpha)",
        ],
    )

    _add_bullets(
        prs,
        "Risque — Justification",
        [
            "KNN : baseline non paramétrique, frontières locales après espace scalé",
            "SVM RBF : séparations non linéaires en espace de features (évent. PCA)",
            "Arbre : règles lisibles ; élagage pour limiter sur-apprentissage",
        ],
    )

    _add_bullets(
        prs,
        "Risque — Complexité",
        [
            "KNN : prédiction O(n) par requête ; stockage des données d'entraînement",
            "SVM : coût d'entraînement plus élevé ; nombre de support vectors variable",
            "Arbre : inférence rapide ; profondeur et feuilles bornent la taille",
        ],
    )

    # --- MODÉLISATION : Avancement ---
    _add_focus_slide(prs, "Modélisation — Prédiction de l'avancement", None)

    _add_bullets(
        prs,
        "Avancement — Modèles utilisés",
        [
            "RandomForestRegressor : GridSearchCV (n_estimators, max_depth, min_samples_split), cv=3, scoring RMSE négatif",
            "LinearRegression : baseline dans Pipeline avec StandardScaler",
        ],
    )

    _add_bullets(
        prs,
        "Avancement — Justification",
        [
            "RF : interactions et non-linéarités ; pas de normalisation requise (commentaire pipeline)",
            "Régression linéaire : simple, interprétable ; hypothèse linéaire stricte",
        ],
    )

    _add_bullets(
        prs,
        "Avancement — Complexité",
        [
            "RF : centaines d'arbres ; entraînement > LR ; prédiction encore rapide",
            "LR : entraînement et inférence très légers une fois features fixées",
        ],
    )

    # --- MODÉLISATION : Priorité ---
    _add_focus_slide(prs, "Modélisation — Recommandation de priorité", None)

    _add_bullets(
        prs,
        "Priorité — Modèles utilisés",
        [
            "RandomForestClassifier : GridSearchCV, scoring f1_macro, cv=3 (rf_priority_pipeline.py)",
            "KNeighborsClassifier : recherche de k sur plage 1–20, baseline (knn_priority_pipeline.py)",
        ],
    )

    _add_bullets(
        prs,
        "Priorité — Justification",
        [
            "RF : agrégation d'arbres ; robuste aux échelles ; optimisé en F1 macro en CV",
            "KNN : baseline comparative ; distances sur features scalées uniquement",
        ],
    )

    _add_bullets(
        prs,
        "Priorité — Complexité",
        [
            "RF : coût d'entraînement modéré ; inférence rapide",
            "KNN : entraînement léger ; prédiction dépend du nombre d'échantillons",
        ],
    )

    # --- ÉVALUATION : Risque ---
    _add_focus_slide(prs, "Évaluation — Segmentation du risque", None)

    _add_bullets(
        prs,
        "Risque — Métriques",
        [
            "Accuracy, précision / rappel / F1 macro, AUC moyenne one-vs-rest (probabilités)",
            "Matrices de confusion, courbes ROC par classe, F1 par classe (figures evaluation/)",
        ],
    )

    _add_table_slide(
        prs,
        "Risque — Résultats (jeu test, risk_knn_svm_dt_evaluation.py)",
        ["Modèle", "Accuracy", "Précision (macro)", "Rappel (macro)", "F1 (macro)", "AUC (OvR)"],
        [
            ["KNN", "0,767", "0,638", "0,619", "0,613", "0,862"],
            ["SVM (RBF)", "0,800", "0,533", "0,598", "0,562", "0,866"],
            ["Decision Tree", "0,683", "0,458", "0,515", "0,476", "0,786"],
        ],
        col_widths=[1.6, 1.35, 1.55, 1.45, 1.45, 1.45],
    )

    _add_bullets(
        prs,
        "Risque — Visualisations",
        ["Insérer matrice de confusion / ROC ici"],
    )

    _add_bullets(
        prs,
        "Risque — Interprétation",
        [
            "SVM : meilleure accuracy (0,80) et AUC légèrement supérieure",
            "KNN : meilleur F1 macro (~0,61) : compromis précision/rappel plus équilibré sur les 3 classes",
            "Arbre : métriques les plus faibles sur ce split ; variance / sur-apprentissage possibles",
            "Écart accuracy vs F1 macro (SVM) : effet possible des classes déséquilibrées",
        ],
        body_size=14,
    )

    # --- ÉVALUATION : Avancement ---
    _add_focus_slide(prs, "Évaluation — Prédiction de l'avancement", None)

    _add_bullets(
        prs,
        "Avancement — Métriques",
        [
            "R², RMSE, MAE, MAPE (%) — progress_rf_lr_evaluation.py / build_evaluation_performance_pptx.py",
            "Graphiques : y_test vs y_pred, résidus (histogramme, boxplot), importances RF",
        ],
    )

    _add_table_slide(
        prs,
        "Avancement — Résultats (jeu test)",
        ["Modèle", "R²", "RMSE", "MAE", "MAPE (%)"],
        [
            ["Régression linéaire", "0,292", "0,283", "0,215", "95,0"],
            ["Random Forest", "0,530", "0,231", "0,151", "76,8"],
        ],
        col_widths=[3.2, 1.5, 1.5, 1.5, 2.0],
    )

    _add_bullets(
        prs,
        "Avancement — Visualisations",
        [
            "Insérer graphiques d'évaluation ici",
            "(scatter y_test vs y_pred, résidus — dossier Prediction avancement/evaluation/figures)",
        ],
    )

    _add_bullets(
        prs,
        "Avancement — Interprétation",
        [
            "Random Forest domine R², RMSE, MAE : non-linéarités et interactions mieux captées",
            "Régression linéaire : baseline utile mais linéarité trop stricte ici",
            "MAPE élevé : valeurs Progress proches de 0 amplifient l'erreur relative (évaluation du dépôt)",
        ],
    )

    # --- ÉVALUATION : Priorité ---
    _add_focus_slide(prs, "Évaluation — Recommandation de priorité", None)

    _add_bullets(
        prs,
        "Priorité — Métriques",
        [
            "Accuracy, précision / rappel / F1 macro et weighted, AUC OvR (evaluation_priorite_output.txt)",
            "Rapports par classe, matrices de confusion, importance des variables (RF)",
        ],
    )

    _add_table_slide(
        prs,
        "Priorité — Résultats (test n=60 stratifié)",
        ["Modèle", "Accuracy", "Précision (macro)", "Rappel (macro)", "F1 (macro)", "AUC (OvR)"],
        [
            ["KNN (baseline)", "0,417", "0,432", "0,448", "0,413", "0,567"],
            ["Random Forest", "0,533", "0,440", "0,414", "0,404", "0,579"],
        ],
        col_widths=[2.2, 1.35, 1.55, 1.45, 1.45, 1.45],
    )

    _add_bullets(
        prs,
        "Priorité — Visualisations",
        ["Insérer matrice de confusion / ROC ici"],
    )

    _add_bullets(
        prs,
        "Priorité — Interprétation",
        [
            "RF : accuracy et rappel High plus élevés (classe majoritaire ~52 % sur le test)",
            "KNN : F1 macro légèrement meilleur ; meilleur rappel Low/Medium, rappel High plus faible",
            "AUC OvR : RF 0,5788 vs KNN 0,5665 (évaluation_priorite_output.txt)",
            "Trade-off : objectif global (RF) vs équité entre classes (KNN) — selon coût métier",
        ],
        body_size=14,
    )

    # --- BENCHMARKING ---
    _add_focus_slide(prs, "Benchmarking", "Vue globale des trois tâches")

    _add_table_slide(
        prs,
        "Comparaison globale (synthèse dépôt)",
        ["Tâche", "Modèles comparés", "Meilleur indicateur clé (test)", "Commentaire court"],
        [
            [
                "Risque (3 classes)",
                "KNN, SVM RBF, Decision Tree",
                "Accuracy max : SVM 0,80 ; F1 macro max : KNN ~0,61",
                "SVM/KNN compétitifs selon la métrique prioritaire",
            ],
            [
                "Avancement (régression)",
                "RF, régression linéaire",
                "R² max : RF 0,530",
                "RF nettement meilleur sur RMSE/MAE",
            ],
            [
                "Priorité (3 classes)",
                "RF, KNN",
                "Accuracy max : RF 0,533 ; F1 macro max : KNN ~0,413",
                "Classes déséquilibrées ; choix selon objectif",
            ],
        ],
        col_widths=[2.0, 2.5, 3.2, 3.0],
        header_pt=10,
        cell_pt=9,
    )

    _add_bullets(
        prs,
        "Choix final — meilleur modèle par tâche",
        [
            "Risque : SVM si l'accuracy globale prime ; KNN si l'équité F1 macro entre classes prime",
            "Avancement : Random Forest (R², RMSE, MAE supérieurs à la régression linéaire)",
            "Priorité : Random Forest pour accuracy / détection High ; KNN pour F1 macro légèrement supérieur",
        ],
        body_size=14,
    )

    _add_bullets(
        prs,
        "Discussion — forces, faiblesses, trade-offs",
        [
            "Forces : prétraitement documenté ; splits reproductibles (random_state=42) ; figures d'évaluation",
            "Faiblesses : jeux modestes ; priorité avec accuracy modérée ; MAPE avancement très élevé",
            "Trade-offs : accuracy vs F1 macro ; simplicité (LR, arbre) vs performance (RF, SVM)",
            "Donnée unique principale : Project-Management-2-enriched.csv (racine ML)",
        ],
        body_size=14,
    )

    # --- CONCLUSION ---
    _add_bullets(
        prs,
        "Conclusion",
        [
            "Trois cas d'usage ML couverts : risque, avancement, priorité — tous évalués sur données projet",
            "Résultats clés : SVM/KNN pour le risque ; RF pour l'avancement ; RF vs KNN pour la priorité",
            "Impact métier aligné avec la valeur Data Science attendue : anticipation, priorisation, KPIs",
            "Poursuite possible : seuils métier, class_weight, features additionnelles (ex. Risk_Level pour priorité)",
        ],
        body_size=14,
    )

    prs.save(OUT_PATH)
    print(f"Enregistré : {OUT_PATH}")


if __name__ == "__main__":
    main()
