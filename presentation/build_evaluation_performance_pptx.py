# -*- coding: utf-8 -*-
"""Génère Evaluation_performances.pptx (métriques, graphiques, interprétations)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ML_ROOT = Path(__file__).resolve().parent.parent
PRESENTATION_DIR = Path(__file__).resolve().parent
OUT_PATH = PRESENTATION_DIR / "Evaluation_performances.pptx"


def _save_presentation(prs: Presentation, path: Path) -> Path:
    """Enregistre le PPTX ; si le fichier cible est ouvert dans PowerPoint, utilise un nom alternatif."""
    try:
        prs.save(path)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}_genere{path.suffix}")
        prs.save(alt)
        print(
            f"Attention : impossible d'écraser {path} (fichier ouvert ?). "
            f"Copie enregistrée : {alt}"
        )
        return alt

# Chemins des figures (générées par les scripts / notebooks d'évaluation)
FIG_RISK = ML_ROOT / "Segmentation du risque" / "evaluation" / "figures"
FIG_PROG = ML_ROOT / "Prediction avancement" / "evaluation" / "figures"
FIG_PRIO = ML_ROOT / "Recommandation priorite" / "evaluation"


def _add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def _add_bullet_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, line in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(18)


def _add_table_slide(
    prs: Presentation,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    footer: str | None = None,
    *,
    body_pt: int = 13,
    header_pt: int = 14,
) -> None:
    layout = prs.slide_layouts[5]  # Title only
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    nrows = len(rows) + 1
    ncols = len(headers)
    left = Inches(0.4)
    top = Inches(1.35)
    width = Inches(12.4)
    height = Inches(min(5.2, 0.35 + 0.30 * nrows))
    table = slide.shapes.add_table(nrows, ncols, left, top, width, height).table
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(header_pt)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = table.cell(i + 1, j)
            c.text = val
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(body_pt)
    if footer:
        box = slide.shapes.add_textbox(Inches(0.4), Inches(6.85), Inches(12.4), Inches(0.55))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = footer
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(64, 64, 64)


def _add_picture_slide(
    prs: Presentation,
    title: str,
    paths: list[Path],
    cols: int = 2,
) -> None:
    layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    existing = [p for p in paths if p.is_file()]
    if not existing:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12), Inches(1))
        box.text_frame.text = "(Fichiers image introuvables — exécutez les scripts d'évaluation.)"
        return
    n = len(existing)
    rows = (n + cols - 1) // cols
    margin_x, margin_y = 0.35, 1.25
    slide_w = 13.333
    usable_w = slide_w - 2 * margin_x
    cell_w = usable_w / cols
    cell_h = (7.5 - margin_y - 0.5) / max(rows, 1)
    for idx, pth in enumerate(existing):
        r, c = divmod(idx, cols)
        x = margin_x + c * cell_w
        y = margin_y + r * cell_h
        # Laisser une petite marge dans la cellule
        slide.shapes.add_picture(
            str(pth),
            Inches(x + 0.05),
            Inches(y + 0.05),
            width=Inches(cell_w - 0.15),
            height=Inches(cell_h - 0.15),
        )


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _add_title_slide(
        prs,
        "Évaluation des performances & benchmarking",
        "Projet ML — gestion de projet (sections 5 et 6)",
    )

    _add_bullet_slide(
        prs,
        "5 — Évaluation des performances",
        [
            "Choix des métriques adaptées (accuracy, F1-score, RMSE, etc.).",
            "Présentation des résultats sous forme de tableaux ou graphiques (matrice de confusion, courbes ROC, etc.).",
            "Interprétation des résultats obtenus (obligatoire).",
        ],
    )

    _add_bullet_slide(
        prs,
        "Choix des métriques — Classification",
        [
            "Accuracy : part de prédictions correctes (simple, mais trompeuse si classes déséquilibrées).",
            "Précision / Rappel / F1 par classe : qualité des positifs prédits et taux de détection.",
            "F1 macro : moyenne non pondérée sur les classes — utile pour l’équité entre Low / Medium / High.",
            "F1 (ou précision/rappel) weighted : pondéré par l’effectif — proche de l’expérience « globale ».",
            "Matrice de confusion : voir quelles classes sont confondues.",
            "ROC + AUC (one-vs-rest) : capacité à séparer chaque classe des autres via scores de probabilité.",
        ],
    )

    _add_bullet_slide(
        prs,
        "Choix des métriques — Régression (avancement)",
        [
            "R² : fraction de variance expliquée (comparaison relative des modèles).",
            "RMSE : erreur quadratique moyenne en unité de la cible — sensible aux grands écarts.",
            "MAE : erreur moyenne en valeur absolue — plus robuste aux outliers extrêmes que le RMSE.",
            "MAPE : erreur relative moyenne (%) — utile pour communiquer, à manier avec prudence si y proche de 0.",
            "Graphiques : y_test vs y_pred, résidus (histogramme, boxplot) pour détecter biais et hétéroscédasticité.",
        ],
    )

    _add_table_slide(
        prs,
        "Résultats — Segmentation du risque (classification, test)",
        ["Modèle", "Accuracy", "Précision (macro)", "Rappel (macro)", "F1 (macro)", "AUC (OvR)"],
        [
            ["KNN", "0,767", "0,638", "0,619", "0,613", "0,862"],
            ["SVM (RBF)", "0,800", "0,533", "0,598", "0,562", "0,866"],
            ["Decision Tree", "0,683", "0,458", "0,515", "0,476", "0,786"],
        ],
        footer="Weighted (aperçu) : SVM meilleur F1 weighted (~0,76) ; KNN ~0,75 ; arbre ~0,64.",
    )

    _add_bullet_slide(
        prs,
        "Interprétation — Segmentation du risque",
        [
            "Le SVM obtient la meilleure accuracy (0,80) et une AUC légèrement supérieure : bonne séparation globale.",
            "Le KNN affiche le meilleur F1 macro (~0,61) : meilleur compromis précision/rappel quand on traite les trois niveaux de risque avec le même poids.",
            "L’arbre de décision est le plus faible sur ces métriques : plus de variance / sur-apprentissage possible malgré l’élagage.",
            "L’écart accuracy vs F1 macro pour le SVM suggère un effet des classes : le modèle peut privilégier la classe majoritaire sur certaines erreurs.",
        ],
    )

    _add_picture_slide(
        prs,
        "Segmentation du risque — Matrices de confusion et ROC (OvR)",
        [FIG_RISK / "confusion_matrices_knn_svm_dt.png", FIG_RISK / "roc_ovr_knn_svm_dt.png"],
        cols=2,
    )
    _add_picture_slide(
        prs,
        "Segmentation du risque — F1 par classe et courbes d’apprentissage",
        [FIG_RISK / "f1_par_classe_knn_svm_dt.png", FIG_RISK / "learning_curves_knn_svm_dt.png"],
        cols=2,
    )

    _add_table_slide(
        prs,
        "Résultats — Prédiction de l’avancement (régression, test)",
        ["Modèle", "R²", "RMSE", "MAE", "MAPE (%)"],
        [
            ["Régression linéaire", "0,292", "0,283", "0,215", "95,0"],
            ["Random Forest", "0,530", "0,231", "0,151", "76,8"],
        ],
        footer="Même split train/test (20 %, random_state=42) que les pipelines du dossier MODEL.",
    )

    _add_bullet_slide(
        prs,
        "Interprétation — Prédiction de l’avancement",
        [
            "La forêt aléatoire domine sur R², RMSE et MAE : elle capte des non-linéarités et des interactions entre variables opérationnelles.",
            "La régression linéaire reste une baseline utile : rapide et interprétable, mais hypothèse de linéarité trop stricte ici.",
            "Le MAPE élevé s’explique en partie par des valeurs d’avancement proches de 0 (erreurs relatives amplifiées).",
            "Les graphiques de résidus permettent de vérifier l’absence de structure systématique résiduelle (surtout pour la baseline).",
        ],
    )

    _add_picture_slide(
        prs,
        "Avancement — Prédictions et résidus",
        [
            FIG_PROG / "scatter_ytest_vs_ypred_rf_lr.png",
            FIG_PROG / "residuals_hist_kde_rf_lr.png",
        ],
        cols=2,
    )
    _add_picture_slide(
        prs,
        "Avancement — Boxplot des résidus et importances (RF)",
        [
            FIG_PROG / "residuals_boxplot_rf_lr.png",
            FIG_PROG / "rf_progress_top3_importance.png",
        ],
        cols=2,
    )
    _add_picture_slide(
        prs,
        "Avancement — Résidus vs valeurs ajustées (régression linéaire)",
        [FIG_PROG / "residuals_vs_fitted_lr.png"],
        cols=1,
    )

    _add_table_slide(
        prs,
        "Résultats — Recommandation de priorité (classification, test)",
        ["Modèle", "Accuracy", "Précision (macro)", "Rappel (macro)", "F1 (macro)", "AUC (OvR)"],
        [
            ["KNN (baseline)", "0,417", "0,432", "0,448", "0,413", "0,567"],
            ["Random Forest", "0,533", "0,440", "0,414", "0,404", "0,579"],
        ],
        footer="Jeu de test n=60 ; classes déséquilibrées (High ≈ 52 % sur le test stratifié).",
    )

    _add_bullet_slide(
        prs,
        "Interprétation — Recommandation de priorité",
        [
            "Le Random Forest maximise l’accuracy et le F1 sur la classe High : comportement attendu avec une classe majoritaire.",
            "Le KNN a un F1 macro légèrement supérieur : meilleur rappel sur Low et Medium grâce à des frontières locales, au prix du rappel sur High.",
            "KNN exige un StandardScaler ajusté sur le train uniquement — sans scaling, les distances seraient dominées par Budget / durée.",
            "Pour la production, le RF est souvent préférable si la détection fiable des tâches très prioritaires et les probabilités (seuils) comptent ; améliorer les minoritaires via class_weight ou données supplémentaires.",
        ],
    )

    _add_picture_slide(
        prs,
        "Priorité — Matrices de confusion",
        [
            FIG_PRIO / "eval_priority_confusion_rf.png",
            FIG_PRIO / "eval_priority_confusion_knn.png",
        ],
        cols=2,
    )
    _add_picture_slide(
        prs,
        "Priorité — F1 par classe et ROC (OvR)",
        [
            FIG_PRIO / "eval_priority_f1_par_classe.png",
            FIG_PRIO / "eval_priority_roc_ovr.png",
        ],
        cols=2,
    )
    _add_picture_slide(
        prs,
        "Priorité — Importance des variables (Random Forest)",
        [FIG_PRIO / "eval_priority_rf_feature_importance.png"],
        cols=1,
    )

    _add_bullet_slide(
        prs,
        "6 — Benchmarking des modèles (obligatoire)",
        [
            "Tableau comparatif des performances sur le jeu de test (métriques alignées avec les scripts d’évaluation).",
            "Justification du modèle retenu pour le déploiement (par cas d’usage : risque, avancement, priorité).",
            "Discussion sur les forces et les faiblesses de chaque approche algorithmique.",
        ],
    )

    _add_table_slide(
        prs,
        "6 — Tableau comparatif des performances (synthèse)",
        [
            "Cas d’usage",
            "Modèles",
            "Indicateur principal",
            "Meilleur résultat",
            "Modèle le plus performant",
        ],
        [
            [
                "Segmentation du risque",
                "KNN, SVM, DT",
                "Accuracy test",
                "0,800",
                "SVM (RBF)",
            ],
            [
                "Segmentation du risque",
                "KNN, SVM, DT",
                "F1 macro (équité classes)",
                "0,613",
                "KNN",
            ],
            [
                "Prédiction de l’avancement",
                "Régression linéaire, RF",
                "R² test",
                "0,530",
                "Random Forest",
            ],
            [
                "Prédiction de l’avancement",
                "Régression linéaire, RF",
                "RMSE test (↓)",
                "0,231",
                "Random Forest",
            ],
            [
                "Recommandation de priorité",
                "KNN, RF",
                "Accuracy test",
                "0,533",
                "Random Forest",
            ],
            [
                "Recommandation de priorité",
                "KNN, RF",
                "AUC ROC (macro OvR)",
                "0,579",
                "Random Forest",
            ],
        ],
        footer="Données : Project-Management-2-enriched (~300 lignes), split 80/20, random_state=42. Les métriques exactes figurent dans les diapos « Résultats » de la section 5.",
        body_pt=12,
        header_pt=12,
    )

    _add_bullet_slide(
        prs,
        "6 — Justification du modèle retenu (déploiement)",
        [
            "Risque : déployer le SVM si l’objectif est l’accuracy globale et une bonne séparation (AUC) ; choisir le KNN si l’on priorise l’équité entre niveaux de risque (F1 macro). L’arbre reste pertinent pour l’interprétabilité (règles) et une latence très faible.",
            "Avancement : Random Forest retenu — meilleur R², RMSE et MAE ; la régression linéaire sert de baseline interprétable mais sous-performante sur ce jeu.",
            "Priorité : Random Forest retenu pour l’API principale — meilleure accuracy et AUC ; le KNN reste utile comme variante (F1 macro légèrement supérieur) et rappelle l’importance du StandardScaler.",
            "Pistes d’amélioration : class_weight / rééquilibrage pour la priorité, validation croisée ou jeu de validation dédié pour éviter tout biais lié au réglage sur le test (notamment risque).",
        ],
    )

    _add_bullet_slide(
        prs,
        "6 — Forces et faiblesses par approche",
        [
            "KNN — Forces : simple, bon F1 macro sur le risque, frontières locales. Faiblesses : sensible à la dimension et au bruit, coût de prédiction qui augmente avec n, nécessite un scaling rigoureux.",
            "SVM (RBF) — Forces : bonnes performances globales sur le risque (accuracy, AUC). Faiblesses : coût d’entraînement plus élevé, sensibilité à C et γ, probabilités à interpréter avec prudence.",
            "Arbre de décision — Forces : modèle explicable, prédiction rapide. Faiblesses : variance élevée, scores plus faibles ici malgré élagage (ccp_alpha).",
            "Random Forest — Forces : robuste, capture interactions et non-linéarités (avancement, priorité), importances de variables. Faiblesses : moins lisible qu’un arbre unique, risque de biais vers la classe majoritaire (priorité).",
            "Régression linéaire — Forces : très rapide, coefficients interprétables après mise à l’échelle. Faiblesses : hypothèse de linéarité limitante pour Progress.",
        ],
    )

    _add_bullet_slide(
        prs,
        "Synthèse (sections 5 et 6)",
        [
            "Section 5 : métriques choisies, tableaux et graphiques (confusion, ROC, résidus) + interprétation obligatoire par cas d’usage.",
            "Section 6 : benchmarking transverse, choix de déploiement argumenté, forces/faiblesses par famille d’algorithmes.",
            "Rappel : lire accuracy avec prudence si classes déséquilibrées ; croiser avec F1 macro, matrices de confusion et courbes ROC.",
        ],
    )

    saved = _save_presentation(prs, OUT_PATH)
    print(f"Enregistré : {saved}")


if __name__ == "__main__":
    main()
